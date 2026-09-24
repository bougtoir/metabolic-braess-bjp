#!/usr/bin/env python3
"""
Generate editable PPTX for SCRR Comment figure.
"""

import os
import importlib.util
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from docx_utils import sanitize_ooxml_package

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_textbox(slide, left, top, width, height, text, font_size=12,
                bold=False, italic=False, color=None, alignment=PP_ALIGN.LEFT):
    if color is None:
        from pptx.dml.color import RGBColor
        color = RGBColor(0, 0, 0)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.font.name = 'Arial'
    p.alignment = alignment
    return txBox


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)

    add_textbox(slide, Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.6),
                "Figure 1. Clonality reduces genetic variance but does not eliminate environmental variance",
                font_size=20, bold=True)

    fig_path = os.path.join(OUTPUT_DIR, "SCRR_Comment_Figure1_Clonal_Environmental_Variance.png")
    if not os.path.exists(fig_path):
        module_name = "scrr_comment"
        script_path = os.path.join(os.path.dirname(__file__), "create_scrr_comment.py")
        spec = importlib.util.spec_from_file_location(module_name, script_path)
        mod = importlib.util.module_from_spec(spec)
        sys.modules[module_name] = mod
        spec.loader.exec_module(mod)
        fig_path = mod.create_figure1()

    max_width = Inches(12.0)
    max_height = Inches(5.6)
    aspect = 4.2 / 13.5
    width = max_width
    height = width * aspect
    if height > max_height:
        height = max_height
        width = height / aspect
    left = (SLIDE_W - width) / 2
    top = Inches(1.0)
    slide.shapes.add_picture(fig_path, left, top, width=width, height=height)

    caption = (
        "Clonality does not eliminate environmental variance. (A) A clonal population has narrow "
        "genetic variance (green) but can still experience wide environmental variance (orange), "
        "especially in research laboratories that do not monitor environmental parameters. "
        "(B) Monthly distribution of PSC-related GEO Series by academic-year start group "
        "(Jan/Feb-start Southern Hemisphere, Mar/Apr-start Japan/Korea, Sep/Oct-start US/EU/CN). "
        "Lines show the proportion of datasets released per month; the dashed grey line is uniform "
        "expectation. Data source: NCBI GEO. "
        "(C) Over time, unmonitored environmental drift can produce outcome variation comparable "
        "to \u2014 or larger than \u2014 genetic effects, even when all cells are isogenic."
    )
    add_textbox(slide, Inches(0.5), Inches(5.0), Inches(12.3), Inches(1.2),
                caption, font_size=11, italic=True, color=None, alignment=PP_ALIGN.LEFT)

    out_path = os.path.join(OUTPUT_DIR, "SCRR_Comment_Figures_InvisibleVariables.pptx")
    prs.save(out_path)
    sanitize_ooxml_package(out_path)
    print(f"Figures PPTX saved to: {out_path}")


if __name__ == "__main__":
    main()
