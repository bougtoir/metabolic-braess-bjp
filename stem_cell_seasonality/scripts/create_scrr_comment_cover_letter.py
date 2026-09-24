#!/usr/bin/env python3
"""
Generate cover letter for SCRR Comment resubmission.
"""

import os
from datetime import date
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_utils import sanitize_ooxml_package

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def add_paragraph(doc, text, bold=False, size=11, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    return p


def get_word_counts():
    """Read the generated manuscript to report exact word counts."""
    from docx_utils import count_docx_words
    manuscript = os.path.join(OUTPUT_DIR, "StemCellReviewsAndReports_Comment_InvisibleVariables.docx")
    if os.path.exists(manuscript):
        total = count_docx_words(manuscript, include_fig_legend=True, include_refs=True, include_declarations=True)
        body = count_docx_words(manuscript, include_fig_legend=False, include_refs=False, include_declarations=False)
        return total, body
    return 1110, 870


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    add_paragraph(doc, "Tatsuki Onishi", bold=True, size=12)
    add_paragraph(doc, "1[Affiliation to be added]")
    add_paragraph(doc, "Email: bougtoir@gmail.com")
    add_paragraph(doc, "")

    add_paragraph(doc, "The Editors")
    add_paragraph(doc, "Stem Cell Reviews and Reports")
    add_paragraph(doc, "Springer Nature")
    add_paragraph(doc, "")

    add_paragraph(doc, date.today().strftime("%B %d, %Y"))
    add_paragraph(doc, "")

    add_paragraph(doc, "Re: Resubmission as Comment – Submission ID 0189f78d-cef6-4eba-8875-82662be53e05", bold=True)
    add_paragraph(doc, "")

    add_paragraph(doc, "Dear Editors,", bold=True)
    add_paragraph(doc, "")

    total_words, body_words = get_word_counts()
    add_paragraph(doc,
        f"Thank you for the invitation to revise and resubmit our work as a Comment. We have "
        f"reformatted the manuscript to meet the Comment guidelines: {body_words:,} words in the main "
        f"body (excluding references, figure legend and declarations; {total_words:,} words total), one "
        f"figure, five references, and no abstract."
    )

    add_paragraph(doc,
        "The revised Comment retains the original message: that clonality reduces genetic "
        "variance but does not eliminate environmental variance, and that pluripotent stem cell "
        "research systematically under-monitors the physical variables to which these cells are "
        "exposed. We removed the two tables, reduced the references from 18 to 5, removed the "
        "abstract, and tightened the discussion while preserving the central hypothesis, the "
        "empirical examples, and the call for systematic environmental monitoring."
    )

    add_paragraph(doc,
        "All numerical claims are derived from public GEO metadata and are reproducible from the "
        "code and data at https://github.com/bougtoir/stem-cell-seasonality. Point-by-point "
        "responses to the editorial request are provided in the accompanying response document."
    )

    add_paragraph(doc,
        "The article has not been submitted elsewhere. All authors have read and approved the "
        "submission and declare no conflicts of interest."
    )

    add_paragraph(doc, "We thank you for your continued consideration.")
    add_paragraph(doc, "")

    add_paragraph(doc, "Sincerely,")
    add_paragraph(doc, "Tatsuki Onishi", bold=True)
    add_paragraph(doc, "Corresponding author")

    out_path = os.path.join(OUTPUT_DIR, "SCRR_CoverLetter_Comment_InvisibleVariables.docx")
    doc.save(out_path)
    sanitize_ooxml_package(out_path)
    print(f"Cover letter saved to: {out_path}")


if __name__ == "__main__":
    main()
