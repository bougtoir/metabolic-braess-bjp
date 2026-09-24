#!/usr/bin/env python3
"""Mark up the differences between two manuscript revisions.

Insertions are shown in red, deletions in blue with a strike-through, so the
changes can be inspected without opening two files side by side. Equations,
figures and tables are left untouched; only text paragraphs are marked up.

Usage:
    python3 scripts/make_tracked_changes.py \
        --old  /path/to/previous.docx \
        --new  manuscripts/TIM_ZeroFree_Manuscript_EN_formatted.docx \
        --out  manuscripts/TIM_ZeroFree_Manuscript_EN_tracked.docx
"""

from __future__ import annotations

import argparse
import difflib
import os
import re
from typing import List

from docx import Document
from docx.shared import RGBColor

INS_COLOR = RGBColor(0xC0, 0x00, 0x00)   # red: inserted text
DEL_COLOR = RGBColor(0x00, 0x00, 0xC0)   # blue: deleted text

TOKEN_RE = re.compile(r"\s+|\w+|[^\w\s]")

# minimum similarity for two paragraphs to be treated as the same paragraph
# rewritten rather than one deleted and one inserted
PAIR_THRESHOLD = 0.6


def tokenize(text: str) -> List[str]:
    return TOKEN_RE.findall(text)


def has_math(par) -> bool:
    xml = par._element.xml
    return "m:oMath" in xml or "<w:drawing" in xml


def paragraph_texts(doc: Document) -> List[str]:
    return [p.text for p in doc.paragraphs]


def clear_runs(par) -> None:
    for run in list(par.runs):
        run._element.getparent().remove(run._element)


def add_run(par, text: str, kind: str, model_run=None):
    run = par.add_run(text)
    if model_run is not None:
        run.font.name = model_run.font.name
        run.font.size = model_run.font.size
        run.bold = model_run.bold
        run.italic = model_run.italic
    if kind == "ins":
        run.font.color.rgb = INS_COLOR
    elif kind == "del":
        run.font.color.rgb = DEL_COLOR
        run.font.strike = True
    return run


def markup_paragraph(par, old_text: str) -> bool:
    """Rewrite a paragraph so that its diff against old_text is colour coded."""
    new_text = par.text
    if new_text == old_text:
        return False
    model_run = par.runs[0] if par.runs else None
    old_tokens, new_tokens = tokenize(old_text), tokenize(new_text)
    matcher = difflib.SequenceMatcher(None, old_tokens, new_tokens, autojunk=False)
    ops = matcher.get_opcodes()
    clear_runs(par)
    for tag, i1, i2, j1, j2 in ops:
        if tag == "equal":
            add_run(par, "".join(new_tokens[j1:j2]), "same", model_run)
        else:
            if tag in ("delete", "replace"):
                add_run(par, "".join(old_tokens[i1:i2]), "del", model_run)
            if tag in ("insert", "replace"):
                add_run(par, "".join(new_tokens[j1:j2]), "ins", model_run)
    return True


def insert_deleted_paragraph(reference_par, text: str) -> None:
    """Insert a fully deleted paragraph before reference_par."""
    par = reference_par.insert_paragraph_before("")
    try:
        par.style = reference_par.style
    except Exception:
        pass
    add_run(par, text, "del")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--old", required=True, help="previous revision (.docx)")
    ap.add_argument("--new", required=True, help="current revision (.docx)")
    ap.add_argument("--out", required=True, help="output marked-up (.docx)")
    args = ap.parse_args()

    old_doc = Document(args.old)
    new_doc = Document(args.new)

    old_texts = [t for t in paragraph_texts(old_doc)]
    new_pars = new_doc.paragraphs
    new_texts = [p.text for p in new_pars]

    matcher = difflib.SequenceMatcher(None, old_texts, new_texts, autojunk=False)

    changed = deleted = 0
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            continue
        if tag == "replace":
            # pair a rewritten paragraph with its old counterpart only when the
            # two are recognisably the same paragraph; unpaired ones are shown
            # as a whole insertion or a whole deletion
            unmatched_old = list(range(i1, i2))
            for j in range(j1, j2):
                par = new_pars[j]
                if has_math(par):
                    # equations and figures are left as they are; consume the
                    # matching old paragraph so it is not reported as deleted
                    for i in list(unmatched_old):
                        if difflib.SequenceMatcher(
                                None, old_texts[i], new_texts[j],
                                autojunk=False).quick_ratio() >= PAIR_THRESHOLD:
                            unmatched_old.remove(i)
                            break
                    continue
                best_i, best_ratio = None, 0.0
                for i in unmatched_old:
                    ratio = difflib.SequenceMatcher(
                        None, old_texts[i], new_texts[j],
                        autojunk=False).quick_ratio()
                    if ratio > best_ratio:
                        best_i, best_ratio = i, ratio
                old_text = ""
                if best_i is not None and best_ratio >= PAIR_THRESHOLD:
                    old_text = old_texts[best_i]
                    unmatched_old.remove(best_i)
                if markup_paragraph(par, old_text):
                    changed += 1
            for i in unmatched_old:
                text = old_texts[i].strip()
                if not text:
                    continue
                anchor = new_pars[min(j2, len(new_pars) - 1)]
                insert_deleted_paragraph(anchor, text)
                deleted += 1
        elif tag == "insert":
            for k in range(j1, j2):
                par = new_pars[k]
                if has_math(par):
                    continue
                if markup_paragraph(par, ""):
                    changed += 1
        elif tag == "delete":
            for k in range(i1, i2):
                text = old_texts[k].strip()
                if not text:
                    continue
                anchor = new_pars[min(j1, len(new_pars) - 1)]
                insert_deleted_paragraph(anchor, text)
                deleted += 1

    new_doc.save(args.out)
    print(f"Marked-up manuscript saved: {os.path.abspath(args.out)}")
    print(f"  paragraphs with inline changes: {changed}")
    print(f"  paragraphs shown as deleted:    {deleted}")


if __name__ == "__main__":
    main()
