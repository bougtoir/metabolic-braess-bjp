#!/usr/bin/env python3
"""
Generate IEEE Sensors Journal supporting submission files:

- highlights_ieee_sensors.docx
- title_page_ieee_sensors.docx
- cover_letter_ieee_sensors.docx
- figures_ieee_sensors.pptx
- graphical_abstract_ieee_sensors.png

All numerical values are loaded from ../results/evaluation_summary.json and
../results/demo_summary.json. Figures are taken from ./ieee_figures/.
"""

import json
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches
from pptx.util import Pt as PptxPt
from pptx.enum.text import PP_ALIGN
from PIL import Image, ImageDraw, ImageFont

SCRIPT_DIR = Path(__file__).resolve().parent
OUT_DIR = SCRIPT_DIR
FIG_DIR = SCRIPT_DIR / 'ieee_figures'
RESULTS_DIR = SCRIPT_DIR.parent / 'results'

TITLE = (
    'Noise removal for dynamic vision sensors: '
    'a physics-informed stochastic resonance framework'
)


def load_metrics():
    """Load metrics shared across the supporting files."""
    with open(RESULTS_DIR / 'evaluation_summary.json') as f:
        summary = json.load(f)
    with open(RESULTS_DIR / 'demo_summary.json') as f:
        demo = json.load(f)
    m = summary['methods']
    return {
        'n_recordings': summary['n_recordings'],
        'n_valid': m['fano_filter']['n_valid'],
        'fano_auc': m['fano_filter']['auc_mean'],
        'fano_spr': m['fano_filter']['spr_mean'],
        'fano_nrr': m['fano_filter']['nrr_mean'],
        'demo_removal': demo['removal_pct'],
        'demo_events_in': demo['events_in'],
        'demo_events_residual': demo['events_residual'],
    }


def build_highlights(metrics):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5

    p = doc.add_paragraph()
    run = p.add_run('Highlights')
    run.font.size = Pt(14)
    run.bold = True

    highlights = [
        'Closed-form optimal noise-removal level is derived for event-based sensors.',
        'A five-parameter pixel model (A5) inverts DVS dark-current statistics.',
        f"Fano-factor filter reaches ROC-AUC {metrics['fano_auc']:.3f} while preserving {metrics['fano_spr']*100:.1f}% of signal.",
        'Six-tier calibration hierarchy closes the loop for operational deployment.',
        f"Open pipeline removes {metrics['demo_removal']}% noise and recovers satellite trajectories from public EBSSA data.",
    ]

    for h in highlights:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(h)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

    out = OUT_DIR / 'highlights_ieee_sensors.docx'
    doc.save(str(out))
    print(f"Highlights saved: {out}")


def build_title_page(metrics):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('TITLE PAGE')
    run.font.size = Pt(14)
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run(TITLE)
    run.font.size = Pt(16)
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run('[Author 1] (ORCID: 0000-0000-0000-0000)')
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('[Author 2] (ORCID: 0000-0000-0000-0000)')
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run('[Affiliations to be added]')
    run.font.size = Pt(12)
    run.italic = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run('Corresponding author:')
    run.bold = True

    p = doc.add_paragraph()
    run = p.add_run('[Name, email, telephone, full postal address]')
    run.font.size = Pt(12)
    run.italic = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run('Manuscript information:')
    run.bold = True

    info = [
        'Article type: Full Research Paper',
        'Target journal: IEEE Sensors Journal (ISSN 1530-437X)',
        f"Evaluated recordings: {metrics['n_valid']} of {metrics['n_recordings']} EBSSA recordings",
        'Number of manuscript figures: 7',
        'Number of manuscript tables: 2',
        'Number of references: 25',
        'Number of pages: 8',
    ]
    for item in info:
        p = doc.add_paragraph()
        run = p.add_run(item)
        run.font.size = Pt(12)

    out = OUT_DIR / 'title_page_ieee_sensors.docx'
    doc.save(str(out))
    print(f"Title page saved: {out}")


def build_cover_letter(metrics):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.5

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    doc.add_paragraph('[Date]')
    doc.add_paragraph('[Corresponding author name and affiliation]')

    p = doc.add_paragraph()
    run = p.add_run(
        'Editor-in-Chief\n'
        'IEEE Sensors Journal\n'
        'IEEE'
    )
    run.font.size = Pt(12)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run('Dear Editor,')
    run.font.size = Pt(12)

    body = [
        (
            f"We submit the enclosed manuscript entitled \"{TITLE}\" for "
            f"consideration as a Full Research Paper in IEEE Sensors Journal."
        ),
        (
            f"The manuscript addresses noise removal for event-based (Dynamic Vision Sensor, DVS) "
            f"imagers operating in low-light astronomical observation. We derive a closed-form "
            f"optimal noise-removal level from a covariate-adjusted stochastic resonance model, "
            f"implement it using a five-parameter pixel noise model, and validate it on the public "
            f"Event-Based Space Situational Awareness (EBSSA) dataset. The proposed Fano-factor "
            f"filter achieves ROC-AUC {metrics['fano_auc']:.3f} while preserving "
            f"{metrics['fano_spr']*100:.1f}% of the signal events, substantially exceeding event-only "
            f"baselines. A proof-of-concept demonstration removes {metrics['demo_removal']}% of "
            f"noise and recovers satellite trajectories."
        ),
        (
            "All data and code needed to reproduce the figures and numerical results are "
            "mirrored at https://github.com/bougtoir/dvs-noise-inverse-ieee-sensors. "
            "The manuscript is 8 pages, contains 7 figures and 2 tables, and follows the IEEE "
            "two-column template."
        ),
        (
            "We confirm that this work is original, has not been published previously, and is "
            "not under consideration elsewhere. All authors have approved the final version and "
            "agree to its submission."
        ),
        "Sincerely,",
        "[Corresponding author signature]",
    ]

    for i, text in enumerate(body):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12 if i > 0 else 6)
        run = p.add_run(text)
        run.font.size = Pt(12)

    out = OUT_DIR / 'cover_letter_ieee_sensors.docx'
    doc.save(str(out))
    print(f"Cover letter saved: {out}")


# -----------------------------------------------------------------------------
# Editable figure PPTX
# -----------------------------------------------------------------------------

FIGURES = [
    {
        'file': 'fig1_schematic.png',
        'title': 'Figure 1',
        'caption': (
            'Conceptual schematic of covariate-adjusted stochastic resonance. '
            '(a) Threshold detector; (b) SR effect; (c) covariate adjustment narrows '
            'residual noise and shifts the operating point.'
        ),
    },
    {
        'file': 'fig2_optimal_rho.png',
        'title': 'Figure 2',
        'caption': (
            'Optimal noise model accuracy \u03c1*. '
            '(a) \u03c1* vs. input noise; yellow: SR regime, blue: excess-noise regime. '
            '(b) Peak SNR gain vs. \u03c3/\u03b8.'
        ),
    },
    {
        'file': 'fig3_pipeline_en.png',
        'title': 'Figure 3',
        'caption': (
            'PI-DC-DVS pipeline: (1) A5 noise model + auxiliary channels; '
            '(2) Bayesian inverse solution; (3) probabilistic thinning; '
            '(4) template-free detection; (5) physical validation.'
        ),
    },
    {
        'file': 'fig4_sr_curves.png',
        'title': 'Figure 4',
        'caption': (
            'Stochastic resonance curves (A/\u03b8 = 0.3). Solid lines: analytical; '
            'squares: Monte Carlo. Covariate adjustment shifts the peak rightward.'
        ),
    },
    {
        'file': 'fig5_detection_and_roc.png',
        'title': 'Figure 5',
        'caption': (
            'Detection and classification performance (A/\u03b8 = 0.4). '
            '(a) P_D; (b) P_FA; (c) ROC curves at \u03c3_n/\u03b8 = 1.5.'
        ),
    },
    {
        'file': 'fig7_evaluation_and_application.png',
        'title': 'Figure 6',
        'caption': (
            'Experimental evaluation on EBSSA. '
            '(a) NRR, (b) SPR, (c) F1, (d) ROC-AUC, '
            '(e) AUC comparison, (f) NRR vs. SPR trade-off.'
        ),
    },
    {
        'file': 'fig8_proof_of_concept.png',
        'title': 'Figure 7',
        'caption': (
            'Proof-of-concept on EBSSA Recording #0. '
            '(a) Raw events; (b) noise rate map; (c) per-event noise probability; '
            '(d) residual events; (e) Fano factor map; (f) temporal rate; (g) per-pixel S/N.'
        ),
    },
]

SUPPLEMENTARY = [
    {
        'file': 'fig6_a5_simulation.png',
        'title': 'Supplementary Figure S1',
        'caption': (
            'A5-based noise rate simulation. '
            '(a) Predicted noise rate [evt/s/pix]; '
            '(b) SNR improvement factor at 90% model accuracy; '
            '(c) SNR vs. temperature at fixed illuminance.'
        ),
    },
]


def _add_textbox(slide, left, top, width, height, text, font_size=PptxPt(12),
                 bold=False, alignment=PP_ALIGN.LEFT):
    tx = slide.shapes.add_textbox(left, top, width, height)
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = font_size
    p.font.bold = bold
    p.alignment = alignment
    return tx


def _build_figure_slide(prs, info):
    blank = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(blank)
    img_path = FIG_DIR / info['file']

    _add_textbox(slide, PptxInches(0.5), PptxInches(0.2),
                 PptxInches(12), PptxInches(0.5),
                 info['title'], font_size=PptxPt(24), bold=True)

    if not img_path.exists():
        _add_textbox(slide, PptxInches(2), PptxInches(3),
                     PptxInches(8), PptxInches(1),
                     f'[MISSING: {img_path}]', font_size=PptxPt(18))
        return slide

    with Image.open(img_path) as img:
        img_w, img_h = img.size
    aspect = img_w / img_h

    max_w = PptxInches(12)
    max_h = PptxInches(5.2)
    if aspect > (max_w / max_h):
        w = max_w
        h = int(w / aspect)
    else:
        h = max_h
        w = int(h * aspect)

    left = int((prs.slide_width - w) / 2)
    top = PptxInches(0.9)
    slide.shapes.add_picture(str(img_path), left, top, w, h)

    cap_top = top + h + PptxInches(0.15)
    _add_textbox(slide, PptxInches(0.5), cap_top,
                 PptxInches(12), PptxInches(1.2),
                 info['caption'], font_size=PptxPt(12))
    return slide


def build_figures_pptx():
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)

    # Title slide for identification
    title_slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_textbox(title_slide, PptxInches(0.5), PptxInches(2.5),
                 PptxInches(12), PptxInches(1.5),
                 TITLE, font_size=PptxPt(32), bold=True, alignment=PP_ALIGN.CENTER)
    _add_textbox(title_slide, PptxInches(0.5), PptxInches(4.2),
                 PptxInches(12), PptxInches(0.5),
                 'IEEE Sensors Journal submission figures',
                 font_size=PptxPt(18), alignment=PP_ALIGN.CENTER)

    for info in FIGURES + SUPPLEMENTARY:
        _build_figure_slide(prs, info)

    out = OUT_DIR / 'figures_ieee_sensors.pptx'
    prs.save(str(out))
    print(f"Figure PPTX saved: {out}")


# -----------------------------------------------------------------------------
# Graphical abstract
# -----------------------------------------------------------------------------

def build_graphical_abstract():
    """Create a single PNG graphical abstract from the conceptual schematic."""
    src = FIG_DIR / 'fig1_schematic.png'
    if not src.exists():
        raise FileNotFoundError(f"Missing source for graphical abstract: {src}")

    with Image.open(src) as img:
        img = img.convert('RGB')

    # Target width; preserve aspect ratio, add a title banner.
    target_w = 1600
    ratio = target_w / img.width
    target_h = int(img.height * ratio)
    img = img.resize((target_w, target_h), Image.Resampling.LANCZOS)

    banner_h = 120
    canvas = Image.new('RGB', (target_w, target_h + banner_h), 'white')
    draw = ImageDraw.Draw(canvas)

    # Try to use a default TTF font; fall back to the default bitmap font.
    try:
        title_font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf', 48)
        sub_font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 28)
    except Exception:
        title_font = ImageFont.load_default()
        sub_font = title_font

    title_text = 'Noise removal for dynamic vision sensors'
    sub_text = 'a physics-informed stochastic resonance framework'

    # Center text in banner
    title_bbox = draw.textbbox((0, 0), title_text, font=title_font)
    sub_bbox = draw.textbbox((0, 0), sub_text, font=sub_font)
    title_x = (target_w - (title_bbox[2] - title_bbox[0])) // 2
    sub_x = (target_w - (sub_bbox[2] - sub_bbox[0])) // 2
    draw.text((title_x, 20), title_text, fill='black', font=title_font)
    draw.text((sub_x, 75), sub_text, fill='black', font=sub_font)

    canvas.paste(img, (0, banner_h))

    out = OUT_DIR / 'graphical_abstract_ieee_sensors.png'
    canvas.save(out, dpi=(300, 300))
    print(f"Graphical abstract saved: {out}")


def main():
    metrics = load_metrics()
    build_highlights(metrics)
    build_title_page(metrics)
    build_cover_letter(metrics)
    build_figures_pptx()
    build_graphical_abstract()
    print('IEEE Sensors Journal supporting files complete.')


if __name__ == '__main__':
    main()
