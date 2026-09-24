from pathlib import Path
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).parent
SUMMARY = ROOT / 'simulation_summary_en.txt'

# Read the latest numerical values from the simulation summary
summary_text = SUMMARY.read_text(encoding='utf-8') if SUMMARY.exists() else ''
lambda0_mean = '0.014'
lambda0_prob = '0.00'
lambda3_mean = '0.917'
lambda3_prob = '0.82'
m = re.search(r'lambda=0\.0:.*final public-will alignment mean=([\d.]+).*P\(final>0\.8\)=([\d.]+)', summary_text, re.S)
if m:
    lambda0_mean = m.group(1)
    lambda0_prob = m.group(2)
m = re.search(r'lambda=3\.0:.*final public-will alignment mean=([\d.]+).*P\(final>0\.8\)=([\d.]+)', summary_text, re.S)
if m:
    lambda3_mean = m.group(1)
    lambda3_prob = m.group(2)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

figures = [
    (ROOT / 'simulation_p1_en.png',
     'Figure 1 Unimodality of the constitutional-cycle loss function L(T)',
     'Left panel: detail near the optimum for T = 2 to 50. Right panel: the same function over the full range T = 2 to 200. The analytical optimum T* and the integer global minimum are indicated for an illustrative parameter set. L(T) is unimodal over 200 random parameter draws and is steeper to the left of the optimum, so erring on the side of a longer cycle is less costly.'),
    (ROOT / 'simulation_p2_en.png',
     'Figure 2 Intertemporal accountability mechanism',
     f'Left: without a trust election (lambda = 0) public-will alignment collapses (mean {lambda0_mean}, P(final>0.8)={lambda0_prob}). Right: with a periodic high-value trust election (lambda = 3) the population converges distributionally to high public-will alignment (mean {lambda3_mean}, P(final>0.8)={lambda3_prob}) (illustrative parameter set; log scale). The parameter lambda is the weight of a future selection event, not a direct reward for the public-will outcome, so the result is a stochastic threshold rather than a mechanical identity.'),
]

for img_path, title, caption in figures:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    slide.shapes.add_picture(str(img_path), Inches(1.5), Inches(1.1), width=Inches(10.333))
    cap_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(12.333), Inches(0.8))
    tf = cap_box.text_frame
    p = tf.paragraphs[0]
    p.text = caption
    p.font.size = Pt(16)
    p.alignment = PP_ALIGN.CENTER

DEST = ROOT / 'figures_en.pptx'
prs.save(DEST)
print(f'Saved {DEST}')
