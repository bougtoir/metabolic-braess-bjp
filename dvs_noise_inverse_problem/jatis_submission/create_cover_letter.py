#!/usr/bin/env python3
"""Generate the JATIS cover letter (cover_letter_jatis.docx).

Numbers are taken from results/*.json via manuscript_numbers.py so the letter
cannot drift from the manuscript.
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt

from manuscript_numbers import auc, load_numbers, scaling

OUT_DIR = Path(__file__).resolve().parent
N = load_numbers()

TITLE = ('A training-free, event-driven Poisson likelihood-ratio detector for faint object events '
         'in sparse neuromorphic space-imaging streams')

PRIOR_SUBMISSIONS = [
    ('Journal of Astronomical Telescopes, Instruments, and Systems', '5 June', '18 June'),
    ('Measurement', '18 June', '19 August'),
    ('Results in Engineering', '20 August', '26 August'),
]


def build_cover_letter():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.3
    style.paragraph_format.space_after = Pt(10)

    doc.add_paragraph('[Date]')
    doc.add_paragraph('Editor-in-Chief\nJournal of Astronomical Telescopes, Instruments, and Systems\nSPIE')
    doc.add_paragraph('Dear Editor,')

    doc.add_paragraph(
        f'We submit the manuscript "{TITLE}" for consideration as a regular paper in the Journal of '
        'Astronomical Telescopes, Instruments, and Systems.'
    )
    doc.add_paragraph(
        'Event cameras are increasingly used for optical space situational awareness, where the event '
        'stream is extremely sparse and dominated by pixel background-activity noise. Existing event '
        'denoising filters output heuristic scores whose thresholds must be tuned on labelled data, and '
        'several are implemented on voxel grids whose cost does not fall as the stream becomes sparser. '
        'We describe a detector that works directly on the asynchronous stream: each pixel\'s background '
        'rate is estimated from its own inter-event intervals, and every event is scored by the Poisson '
        'tail probability of the count in a causal spatiotemporal neighbourhood. The output is a '
        'probability, so the operating point can be set from a false-alarm budget without labels.'
    )
    doc.add_paragraph(
        f'On all {N["n_rec"]} labelled recordings of the public Event-Based Space Situational Awareness '
        f'(EBSSA) dataset, with the published object annotations as ground truth, the detector reaches a '
        f'mean ROC-AUC of {auc(N, "edlr")}, statistically indistinguishable from most published '
        f'unsupervised event filters tuned by the same grouped cross-validation (event bilateral filter '
        f'{auc(N, "bilateral")}) and below the best of them (YNoise {auc(N, "ynoise")}), and its wall-clock '
        f'cost falls with the number of events (exponent {scaling(N, "edlr", "cost_exponent_mean"):.2f}) '
        f'whereas the voxelised form of the same test does not (exponent '
        f'{scaling(N, "plr", "cost_exponent_mean"):.2f}), making the event-driven form '
        f'{scaling(N, "plr", "seconds_sparse") / scaling(N, "edlr", "seconds_sparse"):.0f} times faster '
        'at one percent of the events. We report supervised networks trained on the same labels as an '
        'upper reference and are explicit that the unsupervised detector does not match them nor the '
        'best label-tuned heuristic filter; its value lies in requiring no labels and no manually chosen '
        'threshold, in self-calibration to a stated false-alarm budget, in a null model that follows '
        'directly from the pixel physics of the sensor, and in cost that scales with sparsity, which '
        'suits surveys of unexplored fields, new sensors and on-board processing. A zero-knowledge '
        'analysis shows that a window stated as an expected background count transfers between the two '
        'sensor models of the dataset, and that the same window, handed to YNoise as its time constant '
        f'without labels, raises that filter from {N["hyb"]["ynoise_full"]["default"]["auc"]["mean"]:.3f} '
        f'to {N["hyb"]["ynoise_full"]["handoff"]["auc"]["mean"]:.3f}, close to its label-tuned value on the same recordings.'
    )

    doc.add_paragraph(
        'Disclosure of prior submissions. An earlier and substantially different version of this work '
        'was submitted to, and not accepted by, the following journals:'
    )
    for journal, sub, dec in PRIOR_SUBMISSIONS:
        p = doc.add_paragraph(f'{journal}: submitted {sub}, declined {dec}.')
        p.paragraph_format.left_indent = Pt(24)
        p.paragraph_format.space_after = Pt(2)
    doc.add_paragraph(
        'The present manuscript is a material revision rather than a resubmission of that text. '
        'Specifically: (1) the earlier Fano-factor statistic and its self-referential evaluation have '
        'been removed and are retained only as an ablation; (2) the method has been reformulated as an '
        'event-driven Poisson likelihood-ratio test with per-pixel background-rate estimation; (3) the '
        f'evaluation now covers all {N["n_rec"]} labelled EBSSA recordings with grouped cross-validation '
        'and paired, multiplicity-adjusted tests against the published object annotations; (4) the '
        'comparison includes seven published unsupervised filters, event-driven and voxelised ablations '
        'and two supervised networks; and (5) new analyses quantify computational scaling with event '
        'sparsity and label-free self-calibration to a false-alarm budget. Every number, table and '
        'figure is regenerated by the code in the public repository named in the manuscript.'
    ).paragraph_format.space_before = Pt(6)

    doc.add_paragraph(
        'The manuscript has not been published in a conference proceedings and is not under '
        'consideration elsewhere. All authors have approved the submission and declare no conflicts of '
        'interest. The use of an AI coding assistant in preparing the code and text is disclosed in the '
        'manuscript.'
    )
    doc.add_paragraph('Sincerely,')
    p = doc.add_paragraph('[Corresponding author name, affiliation and e-mail]')
    p.runs[0].italic = True

    out_path = OUT_DIR / 'cover_letter_jatis.docx'
    doc.save(str(out_path))
    print(f'wrote {out_path}')


if __name__ == '__main__':
    build_cover_letter()
