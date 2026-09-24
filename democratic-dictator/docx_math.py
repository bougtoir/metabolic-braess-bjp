import re
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

try:
    from latex2mathml.converter import convert as latex_to_mathml
    from docx_equation import mathml_to_omml
except Exception as exc:  # pragma: no cover - checked at runtime by callers
    latex_to_mathml = None
    mathml_to_omml = None

ITALIC_RE = re.compile(r'\*([^*]+)\*')
MATH_RE = re.compile(r'\$([^$\n]+?)\$')


def add_runs_with_italic(p, text, bold=False, italic=False, font_size=None):
    """Add plain text to a paragraph, italicizing *...* segments."""
    last = 0
    for m in ITALIC_RE.finditer(text):
        start, end = m.start(), m.end()
        if start > last:
            run = p.add_run(text[last:start])
            run.bold = bold
            run.italic = italic
            if font_size:
                run.font.size = Pt(font_size)
        run = p.add_run(m.group(1))
        run.italic = True
        if font_size:
            run.font.size = Pt(font_size)
        last = end
    if last < len(text):
        run = p.add_run(text[last:])
        run.bold = bold
        run.italic = italic
        if font_size:
            run.font.size = Pt(font_size)


def insert_omml(p, latex):
    """Convert LaTeX to OMML and append to a paragraph."""
    if latex_to_mathml is None or mathml_to_omml is None:
        p.add_run(f'[math missing deps: {latex}]')
        return
    try:
        mathml = latex_to_mathml(latex)
        omml = mathml_to_omml(mathml)
        p._element.append(omml)
    except Exception as exc:
        p.add_run(f'[math error: {latex}]')


def add_text_to_para(p, text, bold=False, italic=False, font_size=None):
    """Add text to a paragraph, processing inline $...$ math and *...* italics."""
    last = 0
    for m in MATH_RE.finditer(text):
        start, end = m.start(), m.end()
        if start > last:
            add_runs_with_italic(p, text[last:start], bold=bold, italic=italic, font_size=font_size)
        insert_omml(p, m.group(1))
        last = end
    if last < len(text):
        add_runs_with_italic(p, text[last:], bold=bold, italic=italic, font_size=font_size)


def add_para(doc, text, style_name='Normal', bold=False, center=False, italic=False, font_size=None):
    p = doc.add_paragraph(style=style_name)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text_to_para(p, text, bold=bold, italic=italic, font_size=font_size)
    return p
