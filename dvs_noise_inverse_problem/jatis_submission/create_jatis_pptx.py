#!/usr/bin/env python3
"""Editable English PPTX with one manuscript figure per slide (figures_jatis.pptx).

Widescreen 13.333 x 7.5 in; title at top, image centred and scaled to fit,
caption at the bottom. Captions mirror those in create_jatis_docx.py and
quote numbers from the same result JSONs.
"""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from manuscript_numbers import load_numbers

OUT_DIR = Path(__file__).resolve().parent
N = load_numbers()

FIGURES = [
    ('fig1_method.png', 'Fig. 1 Event-driven Poisson likelihood-ratio detector',
     'Each event is scored from the estimated background rate of its neighbourhood and the number of '
     'events that arrived in that neighbourhood during the preceding window; the tail probability is '
     'the noise probability, so the operating point can be set from a false-alarm budget without labels '
     'and no voxel grid is constructed.'),
    ('fig2_methods_auc.png', 'Fig. 2 Detection performance across methods',
     f'Mean ROC-AUC with standard deviation over the {N["n_rec"]} EBSSA recordings. Dashed line: '
     'proposed detector; dotted line: chance. Colours group the proposed method, published unsupervised '
     'event filters, event-driven ablations, frame-based and label-free learned ablations, and supervised networks.'),
    ('fig3_paired.png', 'Fig. 3 Per-recording paired comparison',
     'Per-recording ROC-AUC of the proposed detector against (a) the event bilateral filter, (b) YNoise, '
     '(c) its voxelised, frame-based form and (d) the supervised MLPF. Circles: 180 x 240-pixel sensor; '
     'squares: 240 x 304-pixel sensor. Points above the diagonal favour the proposed detector.'),
    ('fig4_sparsity_cost.png', 'Fig. 4 Computational cost versus event sparsity',
     '(a) Wall-clock time per recording and (b) time per event as a real EBSSA stream is thinned '
     f'(medians over {N["cost_repeats"]} repeats). Solid: event-driven methods; dashed: frame- or '
     'voxel-based methods. Legend entries give the fitted cost exponent.'),
    ('fig5_calibration.png', 'Fig. 5 Self-calibrated operating points',
     '(a) Median achieved false-positive rate against the requested budget when each score is cut at its '
     'own quantile; dotted line: ideal. (b) Median detection rate at the same budgets. (c) Standardised '
     'partial AUC below false-positive rates of 0.01 and 0.001 (0.5 is chance).'),
    ('fig6_zero_knowledge.png', 'Fig. 6 Zero-knowledge deployment',
     '(a) Mean AUC and median fraction of saturated scores of the rate-adaptive proposal against the '
     'dimensionless window k on the two sensor models; dotted line: deployment value. (b) AUC lost when '
     'the configuration selected on one sensor is applied to the other. (c) Median achieved '
     'false-positive rate on the target sensor with a threshold carried over from the other sensor '
     '(dashed) and recalibrated on the target stream without labels (solid).'),
    ('fig7_hybrid.png', 'Fig. 7 Parameter hand-over and two-stage front end',
     '(a) Mean AUC of the zero-knowledge proposal, of YNoise at its published default and of YNoise with '
     'its time constant set to the rate-adaptive window W handed over from the proposal, pooled and per '
     'sensor model; dotted line: YNoise with label-tuned cross-validation. (b) Mean AUC of the two-stage '
     'score against the fraction of the stream passed by stage 1 (hand-over, default and pseudo-label '
     'parameterisation of stage 2); dashed line: stage 1 alone; black squares (right axis): median '
     'stage-2 time relative to YNoise on the whole stream.'),
    ('fig8_qualitative.png', 'Fig. 8 The same two recordings processed by every stage',
     'Rows: one recording per sensor model ('
     + ', '.join(v['recording_id'] for v in N['qual']['recordings'].values())
     + '), the recording whose stage-1 AUC is closest to the median of that sensor, with its '
     'rate-adaptive window W. Columns: all events; annotated object events; events passed at a '
     f'requested false-alarm rate of {N["qual"]["alpha"]:g} by the zero-knowledge proposal, YNoise at its '
     'published default, the hybrid hand-over (YNoise with dt = W from the proposal), the hybrid two-stage form (proposal screens, YNoise confirms) with f = '
     f'{N["qual"]["screen_fraction"]:g}, and the Fano-factor score. Per-pixel count of passed events on a '
     'logarithmic grey scale; achieved detection and false-positive rates above each panel.'),
]


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for fname, title, caption in FIGURES:
        path = OUT_DIR / fname
        slide = prs.slides.add_slide(blank)
        tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(24)
        p.font.bold = True
        p.alignment = PP_ALIGN.LEFT

        with Image.open(path) as im:
            w, h = im.size
        max_w, max_h = Inches(12.333), Inches(5.0)
        scale = min(max_w / w, max_h / h)
        pw, ph = int(w * scale), int(h * scale)
        left = int((prs.slide_width - pw) / 2)
        top = Inches(1.0) + int((max_h - ph) / 2)
        slide.shapes.add_picture(str(path), left, top, width=pw, height=ph)

        cb = slide.shapes.add_textbox(Inches(0.5), Inches(6.15), Inches(12.333), Inches(1.2))
        cb.text_frame.word_wrap = True
        p = cb.text_frame.paragraphs[0]
        p.text = caption
        p.font.size = Pt(13)
        p.alignment = PP_ALIGN.LEFT
    out = OUT_DIR / 'figures_jatis.pptx'
    prs.save(out)
    print(f'wrote {out} ({len(FIGURES)} slides)')


if __name__ == '__main__':
    main()
