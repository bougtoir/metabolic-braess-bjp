from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_math import add_para, add_text_to_para, insert_omml

ROOT = Path(__file__).parent
SRC = ROOT / 'paper_draft_en_submission.txt'
DEST = ROOT / 'paper_draft_en_submission.docx'

with open(SRC, 'r', encoding='utf-8') as f:
    lines = f.readlines()

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(11)

FIG_RE = re.compile(r'^\[FIGURE:(.+)\]$')
TITLE_LINE_RE = re.compile(r'^(Title Page|Acknowledgments|Statements and Declarations|Abstract|Keywords|References|Data Availability|Figure Legends)$')
SECTION_RE = re.compile(r'^\d+\.\s')
SUBSECTION_RE = re.compile(r'^\d+\.\d+\s')


# Title page: first two non-empty lines are title and subtitle
title_lines = []
body_start = 0
for i, line in enumerate(lines):
    stripped = line.strip()
    if stripped == '' and title_lines:
        body_start = i + 1
        break
    if stripped:
        title_lines.append(stripped)
    else:
        body_start = i + 1

for idx, t in enumerate(title_lines[:2]):
    if idx == 0:
        p = add_para(doc, t, 'Title', font_size=14)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = add_para(doc, t, 'Subtitle', font_size=12)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Body processing
for line in lines[body_start:]:
    s = line.rstrip()
    if s == '' or s == '------------------------------':
        continue
    # Display math
    stripped = s.strip()
    if stripped.startswith('$$') and stripped.endswith('$$'):
        latex = stripped[2:-2].strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        insert_omml(p, latex)
        continue
    # Indented equations / formulas (legacy non-$$)
    if re.match(r'^\s{4,}', line) and not stripped.startswith('where'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        add_text_to_para(p, stripped)
        continue
    # Skip inline figure markers; captions are in Figure Legends at the end
    if FIG_RE.match(s):
        continue
    if SECTION_RE.match(s):
        add_para(doc, s, 'Heading 1')
        continue
    if SUBSECTION_RE.match(s):
        add_para(doc, s, 'Heading 2')
        continue
    if TITLE_LINE_RE.match(s):
        add_para(doc, s, bold=True)
        continue
    if re.match(r'^Figure\s*\d+', s):
        p = add_para(doc, s)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        continue
    if re.match(r'^(Keywords|JEL Classification):', s):
        label, _, rest = s.partition(':')
        p = doc.add_paragraph(style='Normal')
        run = p.add_run(label + ':')
        run.bold = True
        add_text_to_para(p, rest)
        continue
    add_para(doc, s)

doc.save(DEST)
print(f'Saved {DEST}')
