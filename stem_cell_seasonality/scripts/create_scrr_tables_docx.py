#!/usr/bin/env python3
"""
Generate editable separate DOCX for SCRR Hypothesis and Commentary tables.
"""

import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

import importlib.util
import sys

SCRIPT_DIR = os.path.dirname(__file__)
module_path = os.path.join(SCRIPT_DIR, "create_scrr_hypothesis_commentary.py")
spec = importlib.util.spec_from_file_location("scrr_manuscript", module_path)
mod = importlib.util.module_from_spec(spec)
sys.modules["scrr_manuscript"] = mod
spec.loader.exec_module(mod)

OUTPUT_DIR = os.path.join(os.path.dirname(SCRIPT_DIR), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def main():
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Tables for Stem Cell Reviews and Reports submission")
    run.bold = True
    run.font.size = Pt(14)

    p = doc.add_paragraph()
    run = p.add_run("Manuscript: The Invisible Variables: Why Clonal Systems Are Not Immune to Environmental Confounding")
    run.italic = True
    run.font.size = Pt(11)

    # Table 1
    mod.add_gmp_table(doc)

    p = doc.add_paragraph()
    mod.set_paragraph_format(p, space_after=Pt(12))

    # Table 2
    mod.add_iot_table(doc)

    out_path = os.path.join(OUTPUT_DIR, "SCRR_Tables_InvisibleVariables.docx")
    doc.save(out_path)
    print(f"Tables DOCX saved to: {out_path}")


if __name__ == "__main__":
    main()
