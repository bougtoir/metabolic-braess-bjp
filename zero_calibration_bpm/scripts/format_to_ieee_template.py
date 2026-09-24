#!/usr/bin/env python3
"""Reformat the generated TIM manuscript with the IEEE Transactions template.

The template file itself is not redistributed in this repository; download
"Transactions template and instructions on how to create your article" from
IEEE Author Center and pass it with --template.

Usage:
    python3 scripts/format_to_ieee_template.py \
        --template /path/to/Transactions-template....docx \
        [--source manuscripts/TIM_ZeroFree_Manuscript_EN.docx] \
        [--out manuscripts/TIM_ZeroFree_Manuscript_EN_formatted.docx]

What it does:
  * copies the template's styles part into the manuscript package
  * applies template paragraph styles (Title, Authors, Abstract, IndexTerms,
    Heading1/2, Text, Figure Caption, Table Title, References)
  * renumbers section headings as IEEE does (I., II., ... and A., B., ...)
  * sets US Letter page size, template margins and the two-column body
  * suppresses list numbering inherited from the References style so that the
    explicit "[n]" labels are not duplicated
  * appends the author biography section
Equations (OMML), figures and tables are left untouched.
"""
from __future__ import annotations

import argparse
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X",
         "XI", "XII"]
LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

BIOGRAPHY = [
    "Tatsuki Onishi (Member, IEEE) was born in Japan. He received the M.D. "
    "degree from Tohoku University, Sendai, Japan, in 2005, and the Ph.D. "
    "degree in medical informatics from Kyoto University, Kyoto, Japan, in "
    "2021. His major field of study was clinical measurement and biomedical "
    "data science.",
    "He is currently a Project Assistant Professor with the Data Science and "
    "AI Innovation Research Promotion Center, Shiga University, Hikone, Japan. "
    "He was previously a Clinical Researcher with Juntendo University Shizuoka "
    "Hospital and a Visiting Researcher with Healthcare Finland Oy, focusing "
    "on perioperative monitoring and clinical data analysis. His current "
    "research interests include biomedical instrumentation, measurement "
    "uncertainty, clinical data science, and reproducible analysis of "
    "physiological waveforms.",
    "Dr. Onishi is a member of the IEEE and the Japanese Society of "
    "Anesthesiologists. He has served as a reviewer for journals in "
    "measurement and perioperative medicine and received research support "
    "from Shiga University.",
]


def copy_part_xml(src_doc: Document, dst_doc: Document, attr: str) -> None:
    """Replace a package part's XML element in dst with the one from src."""
    src_part = getattr(src_doc.part, attr, None)
    dst_part = getattr(dst_doc.part, attr, None)
    if src_part is None or dst_part is None:
        return
    src_el = src_part.element
    dst_el = dst_part.element
    parent = dst_el.getparent()
    new_el = etree.fromstring(etree.tostring(src_el))
    if parent is None:
        # part root: swap children
        for child in list(dst_el):
            dst_el.remove(child)
        for child in list(new_el):
            dst_el.append(child)
    else:
        parent.replace(dst_el, new_el)
    dst_part._element = new_el


def add_empty_numPr(pPr, num_id: str = "0") -> None:
    """Suppress inherited list numbering from a style by setting numId to 0."""
    for np in pPr.findall(qn("w:numPr")):
        pPr.remove(np)
    np = etree.Element(qn("w:numPr"))
    numId = etree.SubElement(np, qn("w:numId"))
    numId.set(qn("w:val"), num_id)
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is not None:
        idx = list(pPr).index(pStyle)
        pPr.insert(idx + 1, np)
    else:
        pPr.insert(0, np)


def set_style(par, style_name: str, doc: Document) -> bool:
    try:
        par.style = doc.styles[style_name]
        return True
    except KeyError:
        return False


def configure_sections(doc: Document, template: Document) -> None:
    tmpl_sec = template.sections[0]
    for i, sec in enumerate(doc.sections):
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.left_margin = tmpl_sec.left_margin or Inches(0.65)
        sec.right_margin = tmpl_sec.right_margin or Inches(0.65)
        sec.top_margin = tmpl_sec.top_margin or Inches(0.70)
        sec.bottom_margin = tmpl_sec.bottom_margin or Inches(0.70)
        sectPr = sec._sectPr
        for cols in sectPr.findall(qn("w:cols")):
            sectPr.remove(cols)
        cols = etree.SubElement(sectPr, qn("w:cols"))
        if i == 0:  # title page stays single column
            cols.set(qn("w:num"), "1")
            cols.set(qn("w:space"), "720")
        else:
            cols.set(qn("w:num"), "2")
            cols.set(qn("w:space"), "216")
            cols.set(qn("w:equalWidth"), "1")


def strip_frame(doc: Document, style_names) -> None:
    """Remove framePr from styles so multi-line title/author blocks stack."""
    for name in style_names:
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        pPr = st.element.find(qn("w:pPr"))
        if pPr is None:
            continue
        for fr in pPr.findall(qn("w:framePr")):
            pPr.remove(fr)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument(
        "--source",
        default=os.path.join(ROOT, "manuscripts",
                             "TIM_ZeroFree_Manuscript_EN.docx"))
    ap.add_argument(
        "--out",
        default=os.path.join(ROOT, "manuscripts",
                             "TIM_ZeroFree_Manuscript_EN_formatted.docx"))
    args = ap.parse_args()

    template = Document(args.template)
    doc = Document(args.source)

    copy_part_xml(template, doc, "styles")
    strip_frame(doc, ["Title", "Authors"])
    configure_sections(doc, template)

    paras = doc.paragraphs
    h1 = 0
    h2 = 0
    in_refs = False
    abstract_next = False

    for idx, par in enumerate(paras):
        text = par.text.strip()
        style_name = par.style.name if par.style is not None else ""

        if idx == 0 and text:
            set_style(par, "Title", doc)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue

        if abstract_next and text:
            abstract_next = False
            set_style(par, "Abstract", doc)
            if not par.runs:
                par.add_run("")
            new_run = par.runs[0]
            if not text.startswith("Abstract"):
                new_run.text = "Abstract\u2014" + new_run.text
            new_run.bold = True
            new_run.italic = True
            for r in par.runs[1:]:
                r.bold = False
                r.italic = False
            continue

        if style_name.startswith("Heading 1"):
            if text.lower().startswith("abstract"):
                # drop the standalone "Abstract" heading; style the next para
                par._element.getparent().remove(par._element)
                abstract_next = True
                continue
            if text.lower().startswith("reference"):
                in_refs = True
                set_style(par, "Reference Head", doc)
                for run in par.runs:
                    run.text = "REFERENCES"
                continue
            h1 += 1
            h2 = 0
            label = ROMAN[h1 - 1] if h1 <= len(ROMAN) else str(h1)
            clean = re.sub(r"^\s*\d+\.\s*", "", text)
            set_style(par, "Heading 1", doc)
            add_empty_numPr(par._element.get_or_add_pPr())
            for run in par.runs[1:]:
                run.text = ""
            if par.runs:
                par.runs[0].text = f"{label}. {clean.upper()}"
            continue

        if style_name.startswith("Heading 2"):
            h2 += 1
            letter = LETTERS[h2 - 1] if h2 <= len(LETTERS) else str(h2)
            clean = re.sub(r"^\s*\d+\.\d+\.?\s*", "", text)
            set_style(par, "Heading 2", doc)
            add_empty_numPr(par._element.get_or_add_pPr())
            for run in par.runs[1:]:
                run.text = ""
            if par.runs:
                par.runs[0].text = f"{letter}. {clean}"
            continue

        if in_refs and re.match(r"^\[\d+\]", text):
            set_style(par, "References", doc)
            add_empty_numPr(par._element.get_or_add_pPr())
            continue

        if text.startswith("Index Terms"):
            set_style(par, "IndexTerms", doc)
            continue

        if h1 == 0 and (text.startswith("[Author names")
                        or text.startswith("[Affiliations")):
            set_style(par, "Authors", doc)
            continue

        if re.match(r"^(Fig\.|Figure)\s*\d+", text):
            set_style(par, "Figure Caption", doc)
            continue

        if re.match(r"^TABLE|^Table\s*[IVX0-9]+\.", text):
            set_style(par, "Table Title", doc)
            continue

        if text:
            set_style(par, "Text", doc)

    # drop trailing empty paragraphs (and their page breaks) left behind when
    # the tables were moved inline, so the biography follows the references
    for par in reversed(doc.paragraphs):
        if par.text.strip():
            break
        par._element.getparent().remove(par._element)

    # ---------------- author biography ----------------
    doc.add_paragraph()
    for block in BIOGRAPHY:
        p = doc.add_paragraph()
        run = p.add_run(block)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        p.paragraph_format.space_after = Pt(4)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.save(args.out)
    print("Formatted manuscript saved:", args.out)


if __name__ == "__main__":
    main()
