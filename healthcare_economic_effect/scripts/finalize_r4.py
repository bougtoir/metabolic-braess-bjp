"""
Finalize the EHPM R4 manuscript and response/cover.

- Removes reviewer-only / old-version framing.
- Inserts missing spaces before Vancouver-style numeric citations.
- Converts plain-text equations to Word OMML (m:oMath) math zones.
- Replaces 2-byte dashes and minus signs in normal text runs with ASCII equivalents.
- Patches the response letter and cover letter to remove old-version phrasing.
"""
import re
import os
import csv
import json
import copy
from html import escape

from docx import Document
from docx.oxml import parse_xml, OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
DOCX_DIR = os.path.join(ROOT, "output", "docx")

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _is_floatable(s):
    try:
        float(s)
        return True
    except (ValueError, TypeError):
        return False


def _load_ns():
    path = os.path.join(ROOT, "data", "neutral_sustainability.csv")
    with open(path, "r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("iso3") == "JPN":
                return {k: (float(v) if _is_floatable(v) else v) for k, v in row.items()}
    return None


def _load_jcf():
    path = os.path.join(ROOT, "data", "japan_counterfactual.json")
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _load_tempo():
    data = {}
    path = os.path.join(ROOT, "data", "tempo_model_selection.json")
    with open(path, "r", encoding="utf-8") as f:
        data["summary"] = json.load(f)
    csv_path = os.path.join(ROOT, "data", "tempo_model_selection_bycountry.csv")
    with open(csv_path, "r", encoding="utf-8", newline="") as f:
        data["bycountry"] = list(csv.DictReader(f))
    return data


_DATA_CACHE = None


def _get_data():
    global _DATA_CACHE
    if _DATA_CACHE is None:
        _DATA_CACHE = {
            "ns": _load_ns(),
            "jcf": _load_jcf(),
            "tempo": _load_tempo(),
        }
    return _DATA_CACHE

# Tokens that precede figure/table numbers (e.g. 'Figure 1', 'Table 2') -- not reference citations.
NON_CITATION_TOKENS = frozenset({'figure', 'fig', 'figures', 'table', 'tables'})


def _math_zone(inner_xml):
    """Wrap OMML inner XML in a standalone m:oMath element."""
    return f'<m:oMath xmlns:m="{M_NS}" xmlns:w="{W_NS}">{inner_xml}</m:oMath>'


def _mr(text, upright=False):
    """Math run."""
    if upright:
        return f'<m:r><m:rPr><m:sty m:val="text"/></m:rPr><m:t>{escape(text)}</m:t></m:r>'
    return f'<m:r><m:t>{escape(text)}</m:t></m:r>'


def _mvar(text):
    """Italic math variable run."""
    return _mr(text)


def _mtext(text):
    """Upright math text run (e.g. ln, year)."""
    return _mr(text, upright=True)


def _sub(base, subscript):
    return f'<m:sSub><m:e>{base}</m:e><m:sub>{subscript}</m:sub></m:sSub>'


def _frac(num_xml, den_xml):
    return f'<m:f><m:num>{num_xml}</m:num><m:den>{den_xml}</m:den></m:f>'


def _parse_math(xml_str):
    return parse_xml(xml_str)


# --- Specific equation builders ---

def eq_fr_simple():
    """FR = (tau * m) / pf"""
    return _math_zone(
        _mvar("FR") + _mvar("=") +
        _frac(_mvar("τ × m"), _mvar("pf"))
    )


def eq_le_model():
    """LE(t) = alpha + beta ln H(t) + gamma ln GDPpc(t) + epsilon(t)"""
    le_t = _sub(_mvar("LE"), _mvar("t"))
    h_t = _sub(_mvar("H"), _mvar("t"))
    gdppc_t = _sub(_mvar("GDPpc"), _mvar("t"))
    eps_t = _sub(_mvar("ε"), _mvar("t"))
    return _math_zone(
        le_t + _mvar("=") +
        _mvar("α") + _mvar("+") +
        _mvar("β") + _mtext("ln") + _mvar(" ") + h_t + _mvar("+") +
        _mvar("γ") + _mtext("ln") + _mvar(" ") + gdppc_t + _mvar("+") +
        eps_t
    )


def eq_mu_model():
    """mu_H(t) = mu_H0 + mu_H1 * (year - t_0)"""
    mu_h_t = _sub(_mvar("μ") + _mvar("H"), _mvar("t"))
    mu_h0 = _sub(_mvar("μ") + _mvar("H"), _mvar("0"))
    mu_h1 = _sub(_mvar("μ") + _mvar("H"), _mvar("1"))
    t0 = _sub(_mvar("t"), _mvar("0"))
    return _math_zone(
        mu_h_t + _mvar("=") +
        mu_h0 + _mvar("+") + mu_h1 +
        _mvar("×") +
        _mvar("(") + _mtext("year") + _mvar("−") + t0 + _mvar(")")
    )


def eq_m_eff():
    """m_eff = m * (1 - l)"""
    meff = _sub(_mvar("m"), _mvar("eff"))
    return _math_zone(
        meff + _mvar("=") +
        _mvar("m") + _mvar("(") + _mvar("1") + _mvar("−") + _mvar("l") + _mvar(")")
    )


def eq_m_eff_definition():
    """m_eff = m * (1 - l)  where... text not in math"""
    return eq_m_eff()


def eq_fr_output():
    """FR_output = (tau * m) / pf"""
    return _math_zone(
        _sub(_mvar("FR"), _mvar("output")) + _mvar("=") +
        _frac(_mvar("τ × m"), _mvar("pf"))
    )


def eq_fr_va():
    """FR_VA = (tau * m_VA) / pf"""
    return _math_zone(
        _sub(_mvar("FR"), _mvar("VA")) + _mvar("=") +
        _frac(_mvar("τ × ") + _sub(_mvar("m"), _mvar("VA")), _mvar("pf"))
    )


def eq_fr_deficit():
    """FR_deficit = (tau * m_eff * (1-delta)) / pf"""
    meff = _sub(_mvar("m"), _mvar("eff"))
    fr_def = _sub(_mvar("FR"), _mvar("deficit"))
    return _math_zone(
        fr_def + _mvar("=") +
        _frac(_mvar("τ × ") + meff + _mvar(" × ") + _mvar("(") + _mvar("1") + _mvar("−") + _mvar("δ") + _mvar(")"), _mvar("pf"))
    )


def eq_fr_deficit_inline():
    """FR_deficit = tau * m_eff * (1-delta) / pf"""
    return eq_fr_deficit()


def eq_m_eff_inline():
    """m_eff = m(1-l)"""
    return eq_m_eff()


def eq_fr_caption():
    """FR = (tau * m) / pf"""
    return eq_fr_simple()


# --- Inline math token OMML builders ---

# Define the longest / most specific token strings first; ordering controls greedy match.
MATH_TOKENS = [
    # Full inline equations
    ("τ × m / [pf × (1-δ)] = 1.60", _math_zone(
        _frac(_mvar("τ") + _mvar("×") + _mvar("m"),
              _mvar("pf") + _mvar("×") + _mvar("(") + _mtext("1") + _mvar("−") + _mvar("δ") + _mvar(")")) +
        _mvar("=") + _mtext("1.60")
    )),
    ("μ_H0 + μ_H1×[year-t₀]", _math_zone(
        _sub(_mvar("μ"), _mtext("H0")) + _mvar("+") +
        _sub(_mvar("μ"), _mtext("H1")) + _mvar("×") + _mvar("(") +
        _mtext("year") + _mvar("-") + _sub(_mvar("t"), _mtext("0")) + _mvar(")")
    )),
    ("m_eff = m(1 - l)", _math_zone(
        _sub(_mvar("m"), _mtext("eff")) + _mvar("=") +
        _mvar("m") + _mvar("(") + _mtext("1") + _mvar("−") + _mvar("l") + _mvar(")")
    )),
    ("m_eff=m(1 - l)", _math_zone(
        _sub(_mvar("m"), _mtext("eff")) + _mvar("=") +
        _mvar("m") + _mvar("(") + _mtext("1") + _mvar("−") + _mvar("l") + _mvar(")")
    )),
    ("δ_H = 0.10", _math_zone(_sub(_mvar("δ"), _mtext("H")) + _mvar("=") + _mtext("0.10"))),
    ("δ = 0.35", _math_zone(_mvar("δ") + _mvar("=") + _mtext("0.35"))),
    ("μ* = 2.0", _math_zone(_mvar("μ") + _mtext("*") + _mvar("=") + _mtext("2.0"))),
    # Variables / symbols
    ("μ_H0", _math_zone(_sub(_mvar("μ"), _mtext("H0")))),
    ("μ_H1", _math_zone(_sub(_mvar("μ"), _mtext("H1")))),
    ("μ*", _math_zone(_mvar("μ") + _mtext("*"))),
    ("δ_H", _math_zone(_sub(_mvar("δ"), _mtext("H")))),
    ("m_eff", _math_zone(_sub(_mvar("m"), _mtext("eff")))),
    ("t₀", _math_zone(_sub(_mvar("t"), _mtext("0")))),
    ("τ", _math_zone(_mvar("τ"))),
    ("μ", _math_zone(_mvar("μ"))),
    ("δ", _math_zone(_mvar("δ"))),
    ("LE(t)", _math_zone(_sub(_mvar("LE"), _mvar("t")))),
    ("GDPpc(t)", _math_zone(_sub(_mvar("GDPpc"), _mvar("t")))),
    ("H(t)", _math_zone(_sub(_mvar("H"), _mvar("t")))),
    ("m_VA", _math_zone(_sub(_mvar("m"), _mtext("VA")))),
    ("m", _math_zone(_mvar("m"))),
    ("pf", _math_zone(_mvar("pf"))),
    ("l", _math_zone(_mvar("l"))),
    ("×", _math_zone(_mvar("×"))),
    ("₀", _math_zone(_sub(_mvar(""), _mtext("0")))),
]

# Sort by token length (longest first) to avoid partial matches.
MATH_TOKENS.sort(key=lambda x: len(x[0]), reverse=True)


def _text_run(text, rPr=None):
    """Create a w:r element with optional run properties and w:t."""
    r = OxmlElement("w:r")
    if rPr is not None:
        r.append(copy.deepcopy(rPr))
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    return r


def _is_math_boundary(text, pos, token_len):
    """Return True if a math token is not embedded inside a normal word."""
    # Character classes that must not directly precede or follow a token
    # unless the token itself contains that character (e.g. underscores in m_eff).
    alphanum = set("ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789")
    # Preceding character
    if pos > 0:
        before = text[pos - 1]
        if before in alphanum:
            return False
    # Following character (inside the same w:t run)
    after_pos = pos + token_len
    if after_pos < len(text):
        after = text[after_pos]
        if after in alphanum:
            return False
    return True


def _replace_math_tokens_in_element(p_element):
    """Replace math-symbol tokens in w:r/w:t children with inline m:oMath zones."""
    W_R_TAG = f"{{{W_NS}}}r"
    W_RPR_TAG = f"{{{W_NS}}}rPr"
    W_T_TAG = f"{{{W_NS}}}t"

    for child in list(p_element):
        if child.tag != W_R_TAG:
            continue
        t_el = child.find(W_T_TAG)
        if t_el is None or not t_el.text:
            continue
        text = t_el.text
        i = 0
        segments = []
        buf = ""
        while i < len(text):
            matched = False
            for token, omml_xml in MATH_TOKENS:
                if text.startswith(token, i) and _is_math_boundary(text, i, len(token)):
                    if buf:
                        segments.append(("text", buf))
                        buf = ""
                    segments.append(("omml", omml_xml))
                    i += len(token)
                    matched = True
                    break
            if not matched:
                buf += text[i]
                i += 1
        if buf:
            segments.append(("text", buf))
        if not any(kind == "omml" for kind, _ in segments):
            continue

        rPr = child.find(W_RPR_TAG)
        parent = child.getparent()
        idx = list(parent).index(child)
        for kind, value in segments:
            if kind == "text":
                new_node = _text_run(value, rPr)
            else:
                new_node = _parse_math(value)
            parent.insert(idx, new_node)
            idx += 1
        parent.remove(child)


def replace_math_tokens(doc):
    """Convert remaining inline math symbols into Word OMML zones."""
    for p in doc.paragraphs:
        _replace_math_tokens_in_element(p._p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_math_tokens_in_element(p._p)


# --- Helpers for paragraph building ---

def find_para(doc, text, exact=False, start=0):
    for i in range(start, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        if exact:
            if p.text.strip() == text:
                return p
        else:
            if text in p.text:
                return p
    return None


def set_para_text(p, text):
    """Replace paragraph text with plain runs, preserving style/alignment and any existing m:oMath zones."""
    align = p.alignment
    # Remove only text runs; keep paragraph properties and math zones intact.
    for child in list(p._p):
        if child.tag == f"{{{W_NS}}}r":
            p._p.remove(child)
    p.add_run(text)
    p.alignment = align


def replace_substring_in_runs(p, old, new):
    """Replace a substring in the paragraph's w:r runs without touching m:oMath."""
    for run in p.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)


def replace_para_with_omml(p, omml_xml, keep_text_before='', keep_text_after=''):
    """Clear paragraph and insert optional leading text, an OMML zone, and trailing text."""
    align = p.alignment
    p._p.clear()
    if keep_text_before:
        r = p.add_run(keep_text_before)
    p._p.append(_parse_math(omml_xml))
    if keep_text_after:
        r = p.add_run(keep_text_after)
    p.alignment = align


def _add_run_before(p, omml_el, text):
    """Insert a text run immediately before omml_el in p."""
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    r.append(t)
    p._p.insert(p._p.index(omml_el), r)


def _add_run_after(p, omml_el, text):
    """Append a text run after omml_el."""
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    r.append(t)
    # insert after omml_el
    idx = p._p.index(omml_el) + 1
    p._p.insert(idx, r)


def replace_inline_equation(p, pattern, omml_xml, flags=0):
    """Replace the first occurrence of `pattern` in p.text with an OMML zone, leaving surrounding text as runs."""
    full = p.text
    m = re.search(pattern, full, flags)
    if not m:
        return False
    before = full[:m.start()]
    after = full[m.end():]
    align = p.alignment
    p._p.clear()
    if before:
        p.add_run(before)
    omml_el = _parse_math(omml_xml)
    p._p.append(omml_el)
    if after:
        p.add_run(after)
    p.alignment = align
    return True


# --- Manuscript patches ---

def patch_old_version_references(doc):
    # 1. Methods reproducibility paragraph
    p = find_para(doc, "To address the reproducibility concern raised in")
    if p:
        new_text = (
            "The tempo model tests for a constant spending-to-outcome lag in health-capital accumulation. "
            "To ensure reproducibility, the model, the underlying World Bank data, and the model-selection "
            "computation (level RMSE, change RMSE, leave-one-out cross-validation [LOOCV] RMSE, AIC, and BIC) "
            "are all provided in this study's public repository, so the framework can be reproduced and examined "
            "independently. The model treats health expenditure as a stock-building flow and regresses life "
            "expectancy on the log health-capital stock, controlling for income:"
        )
        set_para_text(p, new_text)
        print("Patched old-version reference in tempo intro paragraph.")

    # 2. Tempo model specification paragraph
    p = find_para(doc, "This is a correction to the previous round")
    if p is None:
        p = find_para(doc, "This corrects an earlier version")
    if p:
        t = _get_data()["tempo"]["summary"]
        bycountry = _get_data()["tempo"]["bycountry"]
        n_countries = t["n_countries"]
        period = t["period"]
        mu = t["models"]["M1_constant_lag"]["mu_const_median_yr"]
        mu_H1 = t["models"]["M2_tempo_lag"]["mu_H1_median_yr_per_yr"]
        n_params = t["n_params"]
        vals = [float(r["mu_H1"]) for r in bycountry]
        lower_bound = min(vals)
        upper_bound = max(vals)
        lower_pct = round(sum(1 for v in vals if abs(v - lower_bound) < 0.001) / n_countries * 100)
        upper_pct = round(sum(1 for v in vals if abs(v - upper_bound) < 0.001) / n_countries * 100)
        new_text = (
            f"Three nested variants were compared: M0 (flow-only, no lag; {n_params['M0']} parameters: intercept, ln expenditure, "
            f"ln GDPpc), M1 (constant lag μ*; {n_params['M1']} parameters), and M2 (time-varying lag μ_H0 + μ_H1×[year−t₀]; {n_params['M2']} parameters). "
            f"Lag parameters were estimated by grid search (M1: μ* over 0-19 years in 1-year steps; M2: μ_H0 over 0-18 years "
            f"in 2-year steps and μ_H1 over -0.10 to +0.30 yr/yr in 0.05 steps), refitting the linear model at each grid point and "
            f"selecting the fit that minimized residual sum of squares. Across the {n_countries} countries ({period}) the median constant lag was μ* = {mu:.1f} years. "
            f"The M2 drift parameter μ_H1 was not robustly identified: its country-level estimates were bimodal and dominated by "
            f"the grid boundaries (roughly {lower_pct}% at the lower bound and {upper_pct}% at the upper bound), with a median of {mu_H1:.2f} yr/yr. "
            f"We therefore treat M2 as an exploratory specification rather than as evidence of a stable time-varying drift. "
            f"A positive drift estimate can be obtained if the 2020-2022 COVID-19 mortality shock is included, but that period is not "
            f"representative of the long-term tempo relationship; the {period} sample reported here does not support a positive drift."
        )
        set_para_text(p, new_text)
        print("Patched old-version reference in tempo specification paragraph.")

    # 3. Tempo model performance paragraph
    p = find_para(doc, "This corrects the previous round, in which a positive drift")
    if p is None:
        p = find_para(doc, "This corrects an earlier version, in which a positive drift")
    if p:
        t = _get_data()["tempo"]["summary"]
        m1 = t["models"]["M1_constant_lag"]
        m2 = t["models"]["M2_tempo_lag"]
        kf = t["key_findings"]
        loocv_diff = m1["loocv_rmse_median"] - m2["loocv_rmse_median"]
        mu_H1 = m2["mu_H1_median_yr_per_yr"]
        new_text = (
            f"The time-varying extension (M2 vs M1) did not yield a comparable improvement. The median level RMSE was "
            f"essentially unchanged ({m2['level_rmse_median']:.3f} vs {m1['level_rmse_median']:.3f} years) and, once the additional parameter was penalized, M2 was favored "
            f"over M1 in {kf['M2_beats_M1_aic_pct']}% of countries by AIC and {kf['M2_beats_M1_bic_pct']}% by BIC (median AIC {m2['aic_median']:.1f} vs {m1['aic_median']:.1f}; median BIC {m2['bic_median']:.1f} vs {m1['bic_median']:.1f}); its only "
            f"advantage was on LOOCV RMSE ({kf['M2_beats_M1_loocv_pct']}% of countries, but with a negligible median difference of {loocv_diff:.3f} years). The drift "
            f"parameter μ_H1 was not robustly identified across countries, with a grid-boundary-dominated bimodal distribution and "
            f"a median of {mu_H1:.2f} yr/yr. We therefore find no reliable evidence for a time-varying lag: the robust finding is the existence of "
            f"a constant lag (M1 vs M0), not its time-variation (M2 vs M1). Including the 2020-2022 COVID-19 shock can generate a "
            f"positive drift, but that period is not representative of the long-term tempo relationship. Figure 5 shows the three-layer tempo comparison."
        )
        set_para_text(p, new_text)
        print("Patched old-version reference in tempo performance paragraph.")

    # 4. Results tempo paragraph
    p = find_para(doc, "this corrects the previous round, where a positive drift")
    if p is None:
        p = find_para(doc, "this corrects an earlier version, where a positive drift")
    if p:
        # The current R4 text already contains the desired conclusion; simply remove the old-version tail.
        start = p.text.find("Second,")
        end = p.text.find("this corrects")
        if start != -1 and end != -1:
            new_text = p.text[start:end].strip().rstrip(";,. ")
            if not new_text.endswith("."):
                new_text += "."
            set_para_text(p, new_text)
            print("Patched old-version reference in results tempo paragraph.")

    # Also remove any remaining isolated old-version phrases
    for p in doc.paragraphs:
        replace_substring_in_runs(p, " (see Response to Reviewers)", "")


def polish_language(doc):
    """Minor wording adjustments to reduce AI-like phrasing."""
    substitutions = [
        ("It is important to note that these are output multipliers", "These are output multipliers"),
        ("A critical methodological caveat must be noted: the I-O multipliers", "One caveat is that the I-O multipliers"),
        ("Several further limitations should be noted.", "Several limitations remain."),
        ("A further analytical dimension comes from the Bongaarts-Feeney tempo framework.", "We also draw on the Bongaarts-Feeney tempo framework."),
        ("Japan is also known for its extremely high diagnostic imaging equipment density", "Japan also has a very high diagnostic imaging equipment density"),
        ("Japan faces a uniquely challenging fiscal context", "Japan faces a particularly challenging fiscal context"),
        ("The M2 drift parameter was not robustly identified", "The M2 drift parameter was not reliably identified"),
        ("was not robustly identified", "was not reliably identified"),
        ("produced a substantial and robust improvement", "produced a substantial improvement"),
        ("the robust finding is the existence of a constant lag", "the reliable finding is the existence of a constant lag"),
        ("M1 substantially and robustly outperforms M0", "M1 outperforms M0"),
        ("does not robustly survive", "does not consistently survive"),
        ("This study represents", "This study presents"),
    ]
    for old, new in substitutions:
        for p in doc.paragraphs:
            # Only replace within a single run to avoid splitting OMML runs.
            for run in p.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    print(f"Polished language: {old[:40]}...")
                    break


def patch_equations(doc):
    # 1. Fiscal Return Ratio main equation (centered, standalone)
    p = find_para(doc, "Fiscal Return Ratio = (τ × m) / pf")
    if p:
        replace_para_with_omml(p, eq_fr_simple())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        print("Converted main fiscal return equation to OMML.")

    # 2. LE model equation (centered, standalone)
    p = find_para(doc, "LE(t) = α + β · ln H(t) + γ · ln GDPpc(t) + ε(t)")
    if p:
        replace_para_with_omml(p, eq_le_model())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        print("Converted LE model equation to OMML.")

    # 3. mu_H model equation (centered, standalone)
    p = find_para(doc, "μ_H(t) = μ_H0 + μ_H1 × (year − t₀)")
    if p:
        replace_para_with_omml(p, eq_mu_model())
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        print("Converted mu_H model equation to OMML.")

    # 4. Effective multiplier equation (centered, standalone)
    p = find_para(doc, "m_eff = m × (1 − import leakage rate)")
    if p:
        align = p.alignment
        p._p.clear()
        p._p.append(_parse_math(eq_m_eff()))
        p.add_run(" where l is the import leakage rate.")
        p.alignment = align
        print("Converted effective multiplier equation to OMML.")

    # 5. Table 2 footnote paragraph with three formulas
    p = find_para(doc, "FR (Output) = (τ × m) / pf")
    if p:
        align = p.alignment
        p._p.clear()
        p._p.append(_parse_math(eq_fr_output()))
        p.add_run(", where m is the output multiplier. ")
        p._p.append(_parse_math(eq_fr_va()))
        p.add_run(", where m_VA is the approximate value-added multiplier. ")
        p._p.append(_parse_math(eq_fr_deficit()))
        p.add_run(", where m_eff = m(1 - l), l is the import leakage rate, and δ is the deficit share. Output and VA multipliers bound the demand-side return from above and below; the deficit adjustment is a stylized sustainability boundary, not a claim that bond issuance mechanically reduces contemporaneous tax revenue.")
        p.alignment = align
        print("Converted Table 2 footnote equations to OMML.")

    # 6. Deficit-adjusted paragraph inline equations
    p = find_para(doc, "FR_deficit")
    if p is None:
        p = find_para(doc, "Because roughly")
    if p and "FR_deficit" in p.text:
        align = p.alignment
        text = p.text
        eq1 = "FR_deficit = τ × m_eff × (1−δ) / pf"
        eq2 = "m_eff = m × (1 − import leakage rate)"
        i1 = text.find(eq1)
        i2 = text.find(eq2)
        if i1 != -1 and i2 != -1 and i2 > i1:
            p._p.clear()
            before = text[:i1]
            mid = text[i1 + len(eq1):i2]
            after = text[i2 + len(eq2):]
            if before:
                p.add_run(before)
            p._p.append(_parse_math(eq_fr_deficit_inline()))
            if mid:
                p.add_run(mid)
            p._p.append(_parse_math(eq_m_eff_inline()))
            if after:
                p.add_run(after)
            p.alignment = align
            print("Converted deficit-adjusted paragraph equations to OMML.")

    # 7. Figure 3 caption with equation
    p = find_para(doc, "Fiscal return ratio = (effective tax rate × output multiplier) / public financing share")
    if p:
        align = p.alignment
        p._p.clear()
        p.add_run("Fiscal return ratio: ")
        p._p.append(_parse_math(eq_fr_caption()))
        p.add_run(". Values at or above 1.0 (dashed line) indicate that demand-side tax revenues cover public costs under output-multiplier assumptions.")
        p.alignment = align
        print("Converted Figure 3 caption equation to OMML.")


def _parse_citation_cluster(cluster):
    """Return list of integers for a citation cluster like '8-10' or '1,23'."""
    nums = []
    for part in re.split(r'[,;]', cluster):
        part = part.strip()
        if not part:
            continue
        if '-' in part or '\u2013' in part or '\u2014' in part:
            m = re.match(r'(\d+)\s*[-\u2013\u2014]\s*(\d+)$', part)
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                nums.extend(range(a, b + 1))
        elif part.isdigit():
            nums.append(int(part))
    return nums


def _should_space_before_citation(text, start, cluster):
    """Decide whether a numeric cluster in normal text is a citation needing a preceding space."""
    end = start + len(cluster)
    if end < len(text):
        after = text[end]
        if after.isdigit():
            return False
        if after == '.' and end + 1 < len(text) and text[end + 1].isdigit():
            return False
        # Avoid matching the '1' in '1,000' (thousands separator)
        if after == ',' and end + 3 < len(text) and text[end + 1:end + 4].isdigit():
            return False
        # Reject counts like '12 other', '2 years', '4 million' (space then word)
        if after.isspace():
            i = end + 1
            while i < len(text) and text[i].isspace():
                i += 1
            if i < len(text) and text[i].isalpha():
                return False
    try:
        nums = _parse_citation_cluster(cluster)
    except Exception:
        return False
    if not nums or any(n < 1 or n > 29 for n in nums):
        return False
    if start == 0:
        # A lone number at the start of a run is usually a table count, not a citation.
        if text.strip() == cluster:
            return False
        return False
    before_char = text[start - 1]
    if before_char.isspace():
        return False
    if before_char in ")]'\"\u201d":
        return True
    if before_char.isalpha():
        m = re.search(r'([A-Za-z]+)$', text[:start])
        # Skip M0-M9 model abbreviations (e.g. M1, M2)
        if m and m.group(1) == 'M' and cluster in ['0', '1', '2', '3', '4', '5', '6', '7', '8', '9']:
            return False
        # Skip figure/table references (e.g. Figure 1, Table 2)
        if m and m.group(1).lower().rstrip('.') in NON_CITATION_TOKENS:
            return False
        # Skip subscript digits (e.g. μ_H1, δ_H1)
        if m and len(m.group(1)) == 1 and start >= 2 and text[start - 2] == '_':
            return False
        return True
    if before_char.isdigit():
        # Only split a 4-digit year from a following citation (e.g. 20231,23)
        m = re.search(r'(\d+)$', text[:start])
        if m and len(m.group(1)) == 4:
            idx = start - 4
            if idx == 0 or not text[idx - 1].isalnum():
                return True
        return False
    return False


def normalize_citations(text):
    """Insert a space before citation clusters that are glued to preceding text."""
    pattern = re.compile(r'(?:[1-9]|[12]\d)(?!\d)(?:,(?:[1-9]|[12]\d)(?!\d))*(?:[-\u2013\u2014](?:[1-9]|[12]\d)(?!\d))?')
    positions = [m.start() for m in pattern.finditer(text)
                 if _should_space_before_citation(text, m.start(), m.group(0))]
    out = list(text)
    for pos in sorted(positions, reverse=True):
        out.insert(pos, ' ')
    return ''.join(out)


def _normalize_citations_in_element(p_element):
    """Insert spaces before citations within w:t runs only."""
    for child in p_element:
        if child.tag != f'{{{W_NS}}}r':
            continue
        for t in child.iter(f'{{{W_NS}}}t'):
            if t.text:
                new_text = normalize_citations(t.text)
                if new_text != t.text:
                    t.text = new_text


def normalize_citation_spacing(doc):
    """Add missing spaces before Vancouver-style numeric citations."""
    for p in doc.paragraphs:
        _normalize_citations_in_element(p._p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _normalize_citations_in_element(p._p)


def clean_text_runs(doc):
    """Replace non-ASCII dashes/minus with ASCII equivalents in regular (non-math) runs."""
    for p in doc.paragraphs:
        _clean_runs_in_element(p._p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _clean_runs_in_element(p._p)


def _clean_runs_in_element(p_element):
    """Clean w:r elements only; leave m:oMath elements untouched."""
    for child in p_element:
        if child.tag == f'{{{W_NS}}}r':
            for t in child.iter(f'{{{W_NS}}}t'):
                if t.text:
                    # Replace 2-byte dashes/minus in normal text with ASCII
                    t.text = (t.text
                              .replace('\u2013', '-')   # en dash
                              .replace('\u2014', '--')  # em dash
                              .replace('\u2015', '--')  # horizontal bar
                              .replace('\u2212', '-'))  # minus sign


def patch_response(doc):
    for p in doc.paragraphs:
        t = p.text
        if not t:
            continue
        if "'The previous round' is now 'an earlier version of this analysis'" in t:
            set_para_text(
                p,
                "All reviewer-only phrasing has been removed or rewritten. References to previous rounds and internal cross-references such as '(see Response to Reviewers)' have been deleted from the manuscript."
            )
            print("Patched response comment 6.")


def patch_cover(doc):
    for p in doc.paragraphs:
        t = p.text
        if not t:
            continue
        if "which was previously resubmitted after the de novo rejection" in t:
            new_text = t.split(", which was previously resubmitted after the de novo rejection of EHPM-D-26-00106R2.")[0] + "."
            set_para_text(p, new_text)
            print("Patched cover letter.")


def finalize_manuscript(path):
    doc = Document(path)
    patch_old_version_references(doc)
    # Convert equations to OMML before inserting citation spaces, so math tokens
    # such as μ_H1 or [year-t₀] are not misinterpreted as Vancouver citations.
    patch_equations(doc)
    polish_language(doc)
    clean_text_runs(doc)
    replace_math_tokens(doc)
    normalize_citation_spacing(doc)
    clean_text_runs(doc)
    doc.save(path)
    print(f"Finalized manuscript: {path}")


def finalize_response(path):
    doc = Document(path)
    patch_response(doc)
    clean_text_runs(doc)
    doc.save(path)
    print(f"Finalized response: {path}")


def finalize_cover(path):
    doc = Document(path)
    patch_cover(doc)
    clean_text_runs(doc)
    doc.save(path)
    print(f"Finalized cover: {path}")


def main():
    man = os.path.join(DOCX_DIR, "Healthcare_EHPM_Manuscript_R4.docx")
    resp = os.path.join(DOCX_DIR, "Healthcare_EHPM_ResponseToReviewers_R4.docx")
    cover = os.path.join(DOCX_DIR, "Healthcare_EHPM_CoverLetter_R4.docx")

    if os.path.exists(man):
        finalize_manuscript(man)
    if os.path.exists(resp):
        finalize_response(resp)
    if os.path.exists(cover):
        finalize_cover(cover)


if __name__ == "__main__":
    main()
