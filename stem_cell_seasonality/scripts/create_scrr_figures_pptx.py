#!/usr/bin/env python3
"""
Generate editable PPTX for SCRR Hypothesis and Commentary figures.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_textbox(slide, left, top, width, height, text, font_size=12,
                bold=False, italic=False, color=None, alignment=PP_ALIGN.LEFT):
    if color is None:
        from pptx.dml.color import RGBColor
        color = RGBColor(0, 0, 0)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.font.name = 'Arial'
    p.alignment = alignment
    return txBox


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)

    # Title
    add_textbox(slide, Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.6),
                "Figure 1. Clonality reduces genetic variance but does not eliminate environmental variance",
                font_size=20, bold=True)

    # Image
    fig_path = os.path.join(OUTPUT_DIR, "SCRR_Figure1_Clonal_Environmental_Variance.png")
    if not os.path.exists(fig_path):
        # Regenerate figure by importing the manuscript script function
        import importlib.util
        import sys
        script_path = os.path.join(os.path.dirname(__file__), "create_scrr_hypothesis_commentary.py")
        spec = importlib.util.spec_from_file_location("scrr_manuscript", script_path)
        mod = importlib.util.module_from_spec(spec)
        sys.modules["scrr_manuscript"] = mod
        spec.loader.exec_module(mod)
        fig_path = mod.create_figure1()

    # Scale image to fit slide width with margins (original figure size 10 x 4.5 in)
    max_width = Inches(12.0)
    max_height = Inches(5.6)
    aspect = 4.5 / 10.0
    width = max_width
    height = width * aspect
    if height > max_height:
        height = max_height
        width = height / aspect
    left = (SLIDE_W - width) / 2
    top = Inches(0.95)
    slide.shapes.add_picture(fig_path, left, top, width=width, height=height)

    # Caption
    caption = (
        "Conceptual model of the clonal complacency trap. (A) A clonal population has narrow "
        "genetic variance (green) but can still experience wide environmental variance (orange), "
        "especially in research laboratories that do not monitor environmental parameters. "
        "(B) Over time, unmonitored environmental drift can produce outcome variation comparable "
        "to-or larger than-genetic effects, even when all cells are isogenic."
    )
    add_textbox(slide, Inches(0.5), Inches(6.75), Inches(12.3), Inches(0.7),
                caption, font_size=11, italic=True, color=None, alignment=PP_ALIGN.LEFT)

    out_path = os.path.join(OUTPUT_DIR, "SCRR_Figures_InvisibleVariables.pptx")
    prs.save(out_path)
    print(f"Figures PPTX saved to: {out_path}")


if __name__ == "__main__":
    main()
