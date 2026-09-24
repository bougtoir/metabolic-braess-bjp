#!/usr/bin/env python3
"""
Generate cover letter for Stem Cell Reviews and Reports submission.
"""

import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def add_paragraph(doc, text, bold=False, size=11, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    p.paragraph_format.space_after = Pt(6)
    return p


def main():
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)

    add_paragraph(doc, "Tatsuki Onishi", bold=True, size=12)
    add_paragraph(doc, "[Affiliation to be added]")
    add_paragraph(doc, "Email: bougtoir@gmail.com")
    add_paragraph(doc, "")

    add_paragraph(doc, "The Editors")
    add_paragraph(doc, "Stem Cell Reviews and Reports")
    add_paragraph(doc, "Springer Nature")
    add_paragraph(doc, "")

    from datetime import date
    add_paragraph(doc, date.today().strftime("%B %d, %Y"))
    add_paragraph(doc, "")

    add_paragraph(doc, "Dear Editors,", bold=True)
    add_paragraph(doc, "")

    add_paragraph(doc,
        "We are submitting our manuscript entitled \"The Invisible Variables: Why Clonal Systems "
        "Are Not Immune to Environmental Confounding\" for consideration as a Hypothesis and "
        "Commentary article in Stem Cell Reviews and Reports."
    )

    add_paragraph(doc,
        "The article argues that the pluripotent stem cell field has fallen into a 'clonal "
        "complacency trap': the tacit assumption that eliminating genetic variance removes the "
        "obligation to control environmental variance. We propose the central hypothesis that "
        "clonality reduces genetic variance but does not eliminate environmental variance. We "
        "review evidence that pluripotent stem cells are sensitive to humidity, volatile organic "
        "compounds, ambient light, electromagnetic fields, osmotic stress, and barometric "
        "pressure; summarize the gap between commercial good manufacturing practice facilities "
        "and academic research laboratories; and present an Internet-of-Things-enabled statistical "
        "roadmap for quantifying environmental contributions to differentiation outcomes. We also "
        "explain why retrospective database mining cannot resolve this question because submission "
        "timestamps are confounded by institutional calendars."
    )

    add_paragraph(doc,
        "This Commentary is timely because PSC-derived therapies are entering clinical trials, yet "
        "differentiation outcomes remain variable across laboratories. Identifying and measuring "
        "environmental confounders is a tractable step toward improving reproducibility in both "
        "basic research and clinical-grade manufacturing."
    )

    add_paragraph(doc,
        "The manuscript is approximately 3,000 words, contains one figure and two tables, and "
        "follows Vancouver-style numbered citations. It has not been submitted elsewhere. All "
        "authors have read and approved the submission and declare no conflicts of interest."
    )

    add_paragraph(doc,
        "We thank the Editors and reviewers for their consideration of our work."
    )
    add_paragraph(doc, "")

    add_paragraph(doc, "Sincerely,")
    add_paragraph(doc, "Tatsuki Onishi", bold=True)
    add_paragraph(doc, "Corresponding author")

    out_path = os.path.join(OUTPUT_DIR, "SCRR_CoverLetter_InvisibleVariables.docx")
    doc.save(out_path)
    print(f"Cover letter saved to: {out_path}")


if __name__ == "__main__":
    main()
